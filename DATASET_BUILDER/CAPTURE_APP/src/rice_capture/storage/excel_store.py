from __future__ import annotations

import os
import tempfile
from pathlib import Path
from typing import Any, Iterable

from rice_capture.core.validation import MANUAL_HEADERS, parse_number


class WorkbookStore:
    """Structure-preserving wrapper around an xlsx/xlsm workbook."""

    def __init__(self, path: str | Path):
        self.path = Path(path)
        self.workbook = None
        self.sheet = None
        self.dirty = False

    @staticmethod
    def _openpyxl():
        try:
            import openpyxl
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError as exc:
            raise RuntimeError(
                "Thiếu openpyxl. Hãy chạy setup_capture_app.bat trong CAPTURE_APP."
            ) from exc
        return openpyxl, Alignment, Font, PatternFill

    def load_or_create(self) -> None:
        openpyxl, Alignment, Font, PatternFill = self._openpyxl()
        is_new = not self.path.exists()
        if not is_new:
            keep_vba = self.path.suffix.lower() == ".xlsm"
            self.workbook = openpyxl.load_workbook(self.path, keep_vba=keep_vba)
            self.sheet = self.workbook.active
        else:
            self.workbook = openpyxl.Workbook()
            self.sheet = self.workbook.active
            self.sheet.title = "Sheet1"

        current_headers = [
            str(self.sheet.cell(1, column).value or "").strip()
            for column in range(1, max(self.sheet.max_column, 1) + 1)
        ]
        if not any(current_headers):
            current_headers = []

        changed = False
        added_columns: set[int] = set()
        for header in MANUAL_HEADERS:
            if header not in current_headers:
                current_headers.append(header)
                self.sheet.cell(1, len(current_headers), header)
                added_columns.add(len(current_headers))
                changed = True

        for column, header in enumerate(current_headers, start=1):
            if header and (is_new or column in added_columns):
                cell = self.sheet.cell(1, column)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2F6B4F")
                cell.alignment = Alignment(horizontal="center")
                self.sheet.column_dimensions[cell.column_letter].width = max(16, min(28, len(header) + 3))
        if is_new:
            self.sheet.freeze_panes = "A2"
            self.sheet.auto_filter.ref = self.sheet.dimensions
        self.dirty = changed or is_new

    @property
    def headers(self) -> list[str]:
        self._require_loaded()
        headers = [
            str(self.sheet.cell(1, column).value or "").strip()
            for column in range(1, self.sheet.max_column + 1)
        ]
        while headers and not headers[-1]:
            headers.pop()
        return headers

    def _require_loaded(self) -> None:
        if self.workbook is None or self.sheet is None:
            raise RuntimeError("Workbook chưa được mở.")

    def rows(self) -> list[tuple[int, dict[str, Any]]]:
        self._require_loaded()
        headers = self.headers
        result: list[tuple[int, dict[str, Any]]] = []
        for row_number in range(2, self.sheet.max_row + 1):
            record = {
                header: self.sheet.cell(row_number, column).value
                for column, header in enumerate(headers, start=1)
                if header
            }
            if any(value not in (None, "") for value in record.values()):
                result.append((row_number, record))
        return result

    def contains_sample_id(self, sample_id: str) -> bool:
        wanted = sample_id.strip().upper()
        return any(str(row.get("Sample_ID") or "").strip().upper() == wanted for _, row in self.rows())

    def append(self, record: dict[str, Any]) -> int:
        self._require_loaded()
        sample_id = str(record.get("Sample_ID") or "").strip().upper()
        if not sample_id:
            raise ValueError("Sample_ID không được để trống.")
        if self.contains_sample_id(sample_id):
            raise ValueError(f"Sample_ID {sample_id} đã tồn tại trong workbook.")
        row_number = self.sheet.max_row + 1
        for column, header in enumerate(self.headers, start=1):
            self.sheet.cell(row_number, column, record.get(header))
        self.sheet.auto_filter.ref = self.sheet.dimensions
        self.dirty = True
        return row_number

    def update_cell(self, row_number: int, header: str, value: Any) -> None:
        self._require_loaded()
        if row_number < 2:
            raise ValueError("Không thể sửa hàng tiêu đề.")
        try:
            column = self.headers.index(header) + 1
        except ValueError as exc:
            raise ValueError(f"Không có cột {header}.") from exc
        self.sheet.cell(row_number, column, value)
        self.dirty = True

    def update_manual_value(self, row_number: int, header: str, raw_value: Any) -> None:
        if header == "Sample_ID":
            value = str(raw_value).strip().upper()
            if not value:
                raise ValueError("Sample_ID không được để trống.")
            for other_row, record in self.rows():
                if other_row != row_number and str(record.get("Sample_ID") or "").strip().upper() == value:
                    raise ValueError(f"Sample_ID {value} đã tồn tại.")
            self.update_cell(row_number, header, value)
            return

        if header == "Capture_Timestamp":
            raise ValueError("Capture_Timestamp được máy tính tự tạo khi lưu mẫu.")

        value = parse_number(raw_value, header, integer=header == "Actual_Count")
        if header in {"Container_Height_mm", "Empty_Height_mm"}:
            headers = self.headers
            container = value if header == "Container_Height_mm" else self.sheet.cell(
                row_number, headers.index("Container_Height_mm") + 1
            ).value
            empty = value if header == "Empty_Height_mm" else self.sheet.cell(
                row_number, headers.index("Empty_Height_mm") + 1
            ).value
            if container not in (None, "") and empty not in (None, ""):
                rice_height = float(container) - float(empty)
                if rice_height < 0:
                    raise ValueError("Chiều cao khoảng trống không thể lớn hơn chiều cao ly.")
                self.update_cell(row_number, header, value)
                self.update_cell(row_number, "Rice_Height_mm", round(rice_height, 6))
                return
        self.update_cell(row_number, header, value)

    def delete_rows(self, row_numbers: Iterable[int]) -> None:
        self._require_loaded()
        valid = sorted({int(row) for row in row_numbers if int(row) >= 2}, reverse=True)
        for row in valid:
            self.sheet.delete_rows(row, 1)
        if valid:
            self.sheet.auto_filter.ref = self.sheet.dimensions
            self.dirty = True

    def save(self, path: str | Path | None = None) -> Path:
        self._require_loaded()
        target = Path(path) if path is not None else self.path
        if target.suffix.lower() not in {".xlsx", ".xlsm"}:
            target = target.with_suffix(".xlsx")
        target.parent.mkdir(parents=True, exist_ok=True)
        fd, temp_name = tempfile.mkstemp(prefix=f".{target.stem}_", suffix=target.suffix, dir=target.parent)
        os.close(fd)
        temp_path = Path(temp_name)
        try:
            self.workbook.save(temp_path)
            os.replace(temp_path, target)
        finally:
            if temp_path.exists():
                temp_path.unlink()
        self.path = target
        self.dirty = False
        return target

    def rollback_last_append(self, row_number: int) -> None:
        self._require_loaded()
        if row_number == self.sheet.max_row and row_number >= 2:
            self.sheet.delete_rows(row_number, 1)
            self.dirty = True
