#!/usr/bin/env python3
"""Reconcile legacy raw records and migrate them to a flat M#### image inventory."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import shutil
import sys
from copy import copy
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Sequence, Tuple

import openpyxl

if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if sys.stderr and hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

_SRC_DIR = Path(__file__).resolve().parent.parent / "src"
if str(_SRC_DIR) not in sys.path:
    sys.path.insert(0, str(_SRC_DIR))

from rice_dataset.migration.id_mapping import derive_physical_sample_id, parse_natural_key

VALID_IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff"}
MEASUREMENT_HEADERS = (
    "weight_g",
    "container_height_mm",
    "inner_diameter_mm",
    "empty_height_mm",
    "actual_count",
)


@dataclass
class SourceRow:
    excel_row: int
    sample_id: str
    physical_sample_id: str
    values: List[Any]


@dataclass
class PlanItem:
    sample_id: str
    original_image_id: str
    physical_sample_id: str
    source_image_path: Path
    target_image_filename: str
    source_excel_row: int
    source_excel_id: str
    source_kind: str
    source_values: List[Any]

    def report_dict(self) -> Dict[str, Any]:
        return {
            "Sample_ID": self.sample_id,
            "Original_Image_ID": self.original_image_id,
            "Physical_Sample_ID": self.physical_sample_id,
            "Source_Image_Path": str(self.source_image_path),
            "Target_Image_Filename": self.target_image_filename,
            "Source_Excel_Row": self.source_excel_row,
            "Source_Excel_ID": self.source_excel_id,
            "Migration_Source": self.source_kind,
        }


def _image_files(images_dir: Path) -> List[Path]:
    files = [
        path
        for path in images_dir.rglob("*")
        if path.is_file() and path.suffix.lower() in VALID_IMAGE_EXTENSIONS
    ]
    return sorted(files, key=lambda path: parse_natural_key(path.stem))


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _header_index(headers: Sequence[Any], name: str) -> int | None:
    needle = name.strip().lower()
    for index, value in enumerate(headers):
        if value is not None and str(value).strip().lower() == needle:
            return index
    return None


def _load_source_rows(workbook_path: Path) -> Tuple[List[Any], List[SourceRow], str]:
    workbook = openpyxl.load_workbook(workbook_path, data_only=False)
    worksheet = workbook.active
    headers = [cell.value for cell in worksheet[1]]
    while headers and headers[-1] is None:
        headers.pop()
    sample_index = _header_index(headers, "Sample_ID")
    if sample_index is None:
        raise ValueError("Workbook is missing the Sample_ID column.")

    rows: List[SourceRow] = []
    for excel_row in range(2, worksheet.max_row + 1):
        values = [worksheet.cell(excel_row, col).value for col in range(1, len(headers) + 1)]
        raw_id = values[sample_index]
        if raw_id is None or str(raw_id).strip() == "":
            continue
        sample_id = str(raw_id).strip().upper()
        rows.append(
            SourceRow(
                excel_row=excel_row,
                sample_id=sample_id,
                physical_sample_id=derive_physical_sample_id(sample_id),
                values=values,
            )
        )
    sheet_name = worksheet.title
    workbook.close()
    return headers, rows, sheet_name


def build_reconciliation_plan(
    workbook_path: Path,
    images_dir: Path,
    prefix: str = "M",
    digits: int = 4,
    start: int = 1,
) -> Tuple[List[Any], List[PlanItem], List[SourceRow], List[str], str]:
    headers, rows, sheet_name = _load_source_rows(workbook_path)
    images = _image_files(images_dir)
    errors: List[str] = []

    image_by_stem: Dict[str, List[Path]] = {}
    for path in images:
        image_by_stem.setdefault(path.stem.strip().upper(), []).append(path)
    duplicates = {stem: paths for stem, paths in image_by_stem.items() if len(paths) > 1}
    if duplicates:
        errors.append("Duplicate image stems: " + ", ".join(sorted(duplicates)))

    rows_by_id: Dict[str, List[SourceRow]] = {}
    rows_by_group: Dict[str, List[SourceRow]] = {}
    for row in rows:
        rows_by_id.setdefault(row.sample_id, []).append(row)
        rows_by_group.setdefault(row.physical_sample_id, []).append(row)

    measurement_indices = [
        index
        for name in MEASUREMENT_HEADERS
        for index in [_header_index(headers, name)]
        if index is not None
    ]

    used_excel_rows: set[int] = set()
    plan: List[PlanItem] = []

    for offset, image_path in enumerate(images):
        legacy_id = image_path.stem.strip().upper()
        physical_id = derive_physical_sample_id(legacy_id)
        exact = [row for row in rows_by_id.get(legacy_id, []) if row.excel_row not in used_excel_rows]
        source_kind = "exact_excel_row"

        if exact:
            donor = exact[0]
        else:
            group_rows = rows_by_group.get(physical_id, [])
            if not group_rows:
                errors.append(f"No workbook measurements are available for image {legacy_id}.")
                continue

            signatures = {
                tuple(row.values[index] for index in measurement_indices)
                for row in group_rows
            }
            if len(signatures) != 1:
                errors.append(
                    f"Conflicting manual measurements in physical sample {physical_id}; "
                    f"cannot assign measurements to image {legacy_id}."
                )
                continue

            unused_group_rows = [row for row in group_rows if row.excel_row not in used_excel_rows]
            if unused_group_rows:
                donor = unused_group_rows[0]
                source_kind = "reassigned_same_physical_sample"
            else:
                donor = group_rows[0]
                source_kind = "copied_same_physical_measurements"

        used_excel_rows.add(donor.excel_row)
        new_id = f"{prefix}{start + offset:0{digits}d}"
        plan.append(
            PlanItem(
                sample_id=new_id,
                original_image_id=legacy_id,
                physical_sample_id=physical_id,
                source_image_path=image_path,
                target_image_filename=f"{new_id}{image_path.suffix.lower()}",
                source_excel_row=donor.excel_row,
                source_excel_id=donor.sample_id,
                source_kind=source_kind,
                source_values=list(donor.values),
            )
        )

    excluded_rows = [row for row in rows if row.excel_row not in used_excel_rows]
    if len(plan) != len(images):
        errors.append(f"Planned {len(plan)} rows for {len(images)} images.")

    return headers, plan, excluded_rows, errors, sheet_name


def write_reports(
    report_dir: Path,
    headers: Sequence[Any],
    plan: Sequence[PlanItem],
    excluded_rows: Sequence[SourceRow],
    errors: Sequence[str],
    workbook_path: Path,
    images_dir: Path,
) -> Path:
    report_dir.mkdir(parents=True, exist_ok=True)
    mapping_path = report_dir / "flat_id_mapping.csv"
    with mapping_path.open("w", newline="", encoding="utf-8-sig") as handle:
        fieldnames = list(plan[0].report_dict().keys()) if plan else [
            "Sample_ID", "Original_Image_ID", "Physical_Sample_ID",
            "Source_Image_Path", "Target_Image_Filename",
            "Source_Excel_Row", "Source_Excel_ID", "Migration_Source",
        ]
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(item.report_dict() for item in plan)

    excluded_path = report_dir / "excluded_manual_rows.csv"
    excluded_headers = ["Excel_Row", "Exclusion_Reason"] + [
        str(value) if value is not None else f"Column_{index + 1}"
        for index, value in enumerate(headers)
    ]
    with excluded_path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(excluded_headers)
        for row in excluded_rows:
            writer.writerow([row.excel_row, "No matching raw image after reconciliation"] + row.values)

    source_counts: Dict[str, int] = {}
    for item in plan:
        source_counts[item.source_kind] = source_counts.get(item.source_kind, 0) + 1
    report = {
        "timestamp": datetime.now().isoformat(),
        "workbook_path": str(workbook_path),
        "images_dir": str(images_dir),
        "planned_rows": len(plan),
        "excluded_workbook_rows": len(excluded_rows),
        "source_counts": source_counts,
        "errors": list(errors),
        "apply_ready": not errors,
    }
    report_path = report_dir / "reconciliation_report.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report_path


def _ensure_column(worksheet: Any, name: str) -> int:
    headers = [cell.value for cell in worksheet[1]]
    while headers and headers[-1] is None:
        headers.pop()
    index = _header_index(headers, name)
    if index is not None:
        return index + 1

    new_column = len(headers) + 1
    worksheet.cell(1, new_column, name)
    if new_column > 1:
        source = worksheet.cell(1, new_column - 1)
        target = worksheet.cell(1, new_column)
        if source.has_style:
            target._style = copy(source._style)
        target.font = copy(source.font)
        target.fill = copy(source.fill)
        target.border = copy(source.border)
        target.alignment = copy(source.alignment)
        target.number_format = source.number_format
        target.protection = copy(source.protection)
    return new_column


def stage_workbook(
    workbook_path: Path,
    staged_workbook_path: Path,
    plan: Sequence[PlanItem],
) -> None:
    workbook = openpyxl.load_workbook(workbook_path, data_only=False)
    worksheet = workbook.active
    original_headers = [cell.value for cell in worksheet[1]]
    while original_headers and original_headers[-1] is None:
        original_headers.pop()
    sample_column = (_header_index(original_headers, "Sample_ID") or 0) + 1
    original_id_column = _ensure_column(worksheet, "Original_Sample_ID")
    physical_id_column = _ensure_column(worksheet, "Physical_Sample_ID")
    image_name_column = _ensure_column(worksheet, "Image_Filename")
    source_column = _ensure_column(worksheet, "Migration_Source")
    capture_index = _header_index(original_headers, "Capture_Timestamp")
    rice_height_index = _header_index(original_headers, "Rice_Height_mm")
    container_height_index = _header_index(original_headers, "Container_Height_mm")
    empty_height_index = _header_index(original_headers, "Empty_Height_mm")

    max_column = worksheet.max_column
    for row_index in range(2, worksheet.max_row + 1):
        for column_index in range(1, max_column + 1):
            worksheet.cell(row_index, column_index).value = None

    for output_row, item in enumerate(plan, start=2):
        for column_index, value in enumerate(item.source_values, start=1):
            worksheet.cell(output_row, column_index).value = value
        if (
            rice_height_index is not None
            and container_height_index is not None
            and empty_height_index is not None
        ):
            rice_height = (
                float(item.source_values[container_height_index])
                - float(item.source_values[empty_height_index])
            )
            worksheet.cell(output_row, rice_height_index + 1).value = round(rice_height, 6)


        worksheet.cell(output_row, sample_column).value = item.sample_id
        worksheet.cell(output_row, original_id_column).value = item.original_image_id
        worksheet.cell(output_row, physical_id_column).value = item.physical_sample_id
        worksheet.cell(output_row, image_name_column).value = item.target_image_filename
        worksheet.cell(output_row, source_column).value = item.source_kind

        if capture_index is not None and item.source_kind != "exact_excel_row":
            worksheet.cell(output_row, capture_index + 1).value = None

    try:
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
    except Exception:
        pass

    staged_workbook_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(staged_workbook_path)
    workbook.close()


def stage_images(staged_images_dir: Path, plan: Sequence[PlanItem]) -> None:
    staged_images_dir.mkdir(parents=True, exist_ok=True)
    for item in plan:
        target = staged_images_dir / item.target_image_filename
        shutil.copy2(item.source_image_path, target)
        if _sha256(item.source_image_path) != _sha256(target):
            raise IOError(f"SHA-256 mismatch while staging {item.source_image_path}.")


def verify_staging(
    staged_workbook_path: Path,
    staged_images_dir: Path,
    plan: Sequence[PlanItem],
) -> None:
    workbook = openpyxl.load_workbook(staged_workbook_path, data_only=False)
    worksheet = workbook.active
    headers = [cell.value for cell in worksheet[1]]
    sample_index = _header_index(headers, "Sample_ID")
    image_index = _header_index(headers, "Image_Filename")
    if sample_index is None or image_index is None:
        raise ValueError("Staged workbook is missing Sample_ID or Image_Filename.")

    rows = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if row[sample_index] is not None and str(row[sample_index]).strip():
            rows.append(row)
    workbook.close()

    expected_ids = [item.sample_id for item in plan]
    actual_ids = [str(row[sample_index]).strip().upper() for row in rows]
    if actual_ids != expected_ids:
        raise ValueError("Staged workbook Sample_ID sequence does not match the migration plan.")

    expected_names = {item.target_image_filename for item in plan}
    actual_names = {
        path.name
        for path in staged_images_dir.iterdir()
        if path.is_file() and path.suffix.lower() in VALID_IMAGE_EXTENSIONS
    }
    if actual_names != expected_names:
        raise ValueError("Staged image inventory does not match the migration plan.")

    workbook_names = {str(row[image_index]).strip() for row in rows}
    if workbook_names != expected_names:
        raise ValueError("Staged workbook Image_Filename values do not match staged images.")


def apply_migration(args: argparse.Namespace) -> Dict[str, Any]:
    workbook_path = Path(args.workbook).resolve()
    images_dir = Path(args.images).resolve()
    output_dir = Path(args.output_dir).resolve()

    headers, plan, excluded, errors, sheet_name = build_reconciliation_plan(
        workbook_path, images_dir, args.prefix, args.digits, args.start
    )
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    report_dir = output_dir / f"reconcile_{timestamp}"
    report_path = write_reports(
        report_dir, headers, plan, excluded, errors, workbook_path, images_dir
    )
    if errors:
        raise RuntimeError("Migration plan is blocked:\n - " + "\n - ".join(errors))

    staging_dir = report_dir / "staging"
    staged_workbook = staging_dir / workbook_path.name
    staged_images = staging_dir / "1_Raw_Images"
    stage_workbook(workbook_path, staged_workbook, plan)
    stage_images(staged_images, plan)
    verify_staging(staged_workbook, staged_images, plan)

    backup_dir = output_dir / f"backup_flat_migration_{timestamp}"
    backup_dir.mkdir(parents=True, exist_ok=False)
    backup_workbook = backup_dir / workbook_path.name
    backup_images = backup_dir / "1_Raw_Images"
    shutil.copy2(workbook_path, backup_workbook)

    manifest_path = output_dir / f"flat_migration_manifest_{timestamp}.json"
    manifest = {
        "timestamp": timestamp,
        "status": "STAGED",
        "workbook_path": str(workbook_path),
        "images_dir": str(images_dir),
        "backup_workbook_path": str(backup_workbook),
        "backup_images_dir": str(backup_images),
        "report_path": str(report_path),
        "sheet_name": sheet_name,
        "migrated_samples": len(plan),
        "excluded_workbook_rows": len(excluded),
        "mapping": [item.report_dict() for item in plan],
    }
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    promoted_images = False
    try:
        shutil.move(str(images_dir), str(backup_images))
        shutil.move(str(staged_images), str(images_dir))
        promoted_images = True
        shutil.copy2(staged_workbook, workbook_path)
    except Exception:
        shutil.copy2(backup_workbook, workbook_path)
        if promoted_images and images_dir.exists():
            shutil.rmtree(images_dir)
        if backup_images.exists() and not images_dir.exists():
            shutil.move(str(backup_images), str(images_dir))
        raise

    manifest["status"] = "APPLIED"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return {
        "status": "APPLIED",
        "migrated_samples": len(plan),
        "excluded_workbook_rows": len(excluded),
        "manifest_path": str(manifest_path),
        "backup_dir": str(backup_dir),
        "report_path": str(report_path),
    }


def rollback(manifest_file: Path) -> Dict[str, str]:
    manifest = json.loads(manifest_file.read_text(encoding="utf-8"))
    workbook_path = Path(manifest["workbook_path"]).resolve()
    images_dir = Path(manifest["images_dir"]).resolve()
    backup_workbook = Path(manifest["backup_workbook_path"]).resolve()
    backup_images = Path(manifest["backup_images_dir"]).resolve()

    if images_dir.name != "1_Raw_Images" or len(images_dir.parts) < 3:
        raise ValueError(f"Unsafe images directory in manifest: {images_dir}")
    if not backup_workbook.exists() or not backup_images.exists():
        raise FileNotFoundError("Rollback backup is incomplete.")

    shutil.copy2(backup_workbook, workbook_path)
    if images_dir.exists():
        shutil.rmtree(images_dir)
    shutil.copytree(backup_images, images_dir)
    return {"status": "ROLLED_BACK", "workbook_path": str(workbook_path), "images_dir": str(images_dir)}


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Reconcile legacy rows against existing images and migrate to flat M#### naming.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument("--workbook", default="DATASET_BUILDER/2_Manual_Records/manual_data.xlsx")
    parser.add_argument("--images", default="DATASET_BUILDER/1_Raw_Images")
    parser.add_argument("--output-dir", default="DATASET_BUILDER/migrations")
    parser.add_argument("--prefix", default="M")
    parser.add_argument("--digits", type=int, default=4)
    parser.add_argument("--start", type=int, default=1)
    mode = parser.add_mutually_exclusive_group()
    mode.add_argument("--dry-run", action="store_true")
    mode.add_argument("--apply", action="store_true")
    mode.add_argument("--rollback", type=str)
    args = parser.parse_args()

    if args.rollback:
        result = rollback(Path(args.rollback))
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0

    workbook_path = Path(args.workbook).resolve()
    images_dir = Path(args.images).resolve()
    output_dir = Path(args.output_dir).resolve()

    if not workbook_path.exists():
        raise FileNotFoundError(f"Workbook not found: {workbook_path}")
    if not images_dir.exists():
        raise FileNotFoundError(f"Images directory not found: {images_dir}")
    if args.digits < 1 or args.start < 1:
        raise ValueError("--digits and --start must be positive integers.")

    if args.apply:
        result = apply_migration(args)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0

    headers, plan, excluded, errors, _ = build_reconciliation_plan(
        workbook_path, images_dir, args.prefix, args.digits, args.start
    )
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    report_dir = output_dir / f"reconcile_dry_run_{timestamp}"
    report_path = write_reports(
        report_dir, headers, plan, excluded, errors, workbook_path, images_dir
    )
    summary = {
        "status": "DRY_RUN_READY" if not errors else "DRY_RUN_BLOCKED",
        "planned_rows": len(plan),
        "excluded_workbook_rows": len(excluded),
        "report_path": str(report_path),
        "errors": errors,
    }
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0 if not errors else 1


if __name__ == "__main__":
    sys.exit(main())
